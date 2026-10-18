"""Tests for the spreadsheet import (projectpulse/importer.py) and the transactional load.

Run:  pytest -q tests/test_importer.py
"""
from __future__ import annotations

import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from projectpulse import importer, metrics
from projectpulse.data_generator import generate_dataset
from projectpulse.database import get_engine, import_history, prepare, read_all, record_import, replace_all


@pytest.fixture(scope="module")
def source() -> dict[str, pd.DataFrame]:
    return generate_dataset()


@pytest.fixture(scope="module")
def books(source) -> dict[str, bytes]:
    return importer.write_workbooks(source)


@pytest.fixture(scope="module")
def result(books) -> importer.ImportResult:
    return importer.validate_files(list(books.items()))


def _edit(books: dict[str, bytes], book: str, sheet: str, edits: dict[tuple[int, str], object]) -> list:
    """Copy of ``books`` with cells changed: {(spreadsheet_row, header): value}."""
    wb = load_workbook(io.BytesIO(books[book]))
    ws = wb[sheet]
    col = {c.value: c.column for c in ws[1]}
    for (row, header), value in edits.items():
        ws.cell(row, col[header]).value = value
    buf = io.BytesIO()
    wb.save(buf)
    return [(k, buf.getvalue() if k == book else v) for k, v in books.items()]


def _problems(res: importer.ImportResult, level: str = "error") -> pd.DataFrame:
    p = res.problem_frame()
    return p[p.Level == level.title()]


# --------------------------------------------------------------------------- round trip
def test_round_trip_is_lossless(source, result):
    assert result.ok and not result.problems
    a, b = prepare(source), prepare(result.tables)
    for name in a:
        assert len(a[name]) == len(b[name]), name
    pa = metrics.project_metrics(a, "2026-09-30")
    pb = metrics.project_metrics(b, "2026-09-30")
    assert np.allclose(pa.health_score, pb.health_score)
    assert pa.rag.tolist() == pb.rag.tolist()


def test_templates_have_instructions_and_headers():
    blank = importer.write_workbooks()
    assert set(blank) == {"jira_export.xlsx", "pmo_portfolio.xlsx", "hr_resources.xlsx"}
    wb = load_workbook(io.BytesIO(blank["jira_export.xlsx"]))
    assert wb.sheetnames == ["Instructions", "Issues", "Sprints"]
    assert [c.value for c in wb["Issues"][1]][:3] == ["Issue Key", "Summary", "Issue Type"]


# --------------------------------------------------------------------------- validation
def test_bad_rows_are_reported_with_row_numbers(books):
    files = _edit(books, "jira_export.xlsx", "Issues", {
        (2, "Story Points"): 4, (3, "Assignee"): "Jane Nobody", (5, "Status"): "Waiting for customer"})
    res = importer.validate_files(files)
    errs = _problems(res)
    assert not res.ok and res.tables is None
    assert set(zip(errs.Row, errs.Column)) >= {(2, "Story Points"), (3, "Assignee"), (5, "Status")}


def test_done_without_resolved_date_is_an_error(books, source):
    t = source["tasks"]
    row = int(np.flatnonzero(t.status.to_numpy() == "Done")[0]) + 2
    res = importer.validate_files(_edit(books, "jira_export.xlsx", "Issues", {(row, "Resolved"): None}))
    errs = _problems(res)
    assert ((errs.Row == row) & (errs.Column == "Resolved")).any()


def test_unknown_project_manager(books):
    res = importer.validate_files(_edit(books, "pmo_portfolio.xlsx", "Projects",
                                        {(2, "Project Manager"): "Nobody Here"}))
    assert "Project Manager" in set(_problems(res).Column)


def test_missing_required_sheet():
    blank = importer.write_workbooks()
    res = importer.validate_files([("pmo_portfolio.xlsx", blank["pmo_portfolio.xlsx"])])
    missing = set(_problems(res).Sheet)
    assert {"Issues", "Sprints", "Employees"} <= missing


def test_jira_vocabulary_is_accepted(books, result):
    files = _edit(books, "jira_export.xlsx", "Issues", {
        (2, "Status"): "Closed", (2, "Issue Type"): "User Story", (2, "Priority"): "Highest"})
    res = importer.validate_files(files)
    assert res.ok
    t = res.tables["tasks"].iloc[0]
    assert (t.issue_type, t.priority) == ("Story", "Critical")


def test_csv_sheet_detected_by_headers(books):
    sheets = importer.tables_to_sheets(importer.validate_files(list(books.items())).tables)
    csv = sheets["Issues"].to_csv(index=False).encode()
    others = [(k, v) for k, v in books.items()]
    wb = load_workbook(io.BytesIO(books["jira_export.xlsx"]))
    del wb["Issues"]
    buf = io.BytesIO()
    wb.save(buf)
    files = [(k, buf.getvalue() if k == "jira_export.xlsx" else v) for k, v in others]
    res = importer.validate_files(files + [("export-2026-09-30.csv", csv)])
    assert res.ok, res.problem_frame().head()
    assert res.sheets_found["Issues"] == "export-2026-09-30.csv"


def test_percent_columns_accept_fractions_and_numbers(books):
    res = importer.validate_files(_edit(books, "pmo_portfolio.xlsx", "Risks",
                                        {(r, "Mitigation %"): 40 for r in range(2, 302)}))
    assert res.ok
    assert (res.tables["risks"].mitigation_progress == 40).all()


def test_missing_peer_reviews_rescale_performance(books, source):
    files = [(k, v) for k, v in books.items() if k != "hr_resources.xlsx"]
    wb = load_workbook(io.BytesIO(books["hr_resources.xlsx"]))
    del wb["Peer Reviews"]
    buf = io.BytesIO()
    wb.save(buf)
    res = importer.validate_files(files + [("hr_resources.xlsx", buf.getvalue())])
    assert res.ok
    assert len(_problems(res, "warning"))
    d = prepare(res.tables)
    perf = metrics.employee_performance(d["tasks"], d["performance"], d["sprints"], d["employees"])
    assert perf.performance_score.between(0, 100).all()


def test_archive_optional_when_history_exists(books, source):
    wb = load_workbook(io.BytesIO(books["pmo_portfolio.xlsx"]))
    del wb["Archive"]
    buf = io.BytesIO()
    wb.save(buf)
    files = [(k, buf.getvalue() if k == "pmo_portfolio.xlsx" else v) for k, v in books.items()]
    assert not importer.validate_files(files).ok  # nothing to keep on a fresh database
    kept = importer.validate_files(files, source["project_history"])
    assert kept.ok and len(kept.tables["project_history"]) == len(source["project_history"])


# --------------------------------------------------------------------------- database
def test_replace_all_and_import_log(tmp_path, result):
    engine = get_engine(f"sqlite:///{tmp_path / 'import.db'}")
    counts = replace_all(engine, result.tables)
    record_import(engine, "spreadsheet", counts, ["a.xlsx"], "2026-09-30")
    replace_all(engine, result.tables)  # second load keeps the log
    record_import(engine, "spreadsheet", counts, ["b.xlsx"], "2026-09-30")
    back = read_all(engine)
    assert len(back["tasks"]) == 10_000
    assert back["performance"].peer_review_score.notna().all()
    log = import_history(engine)
    assert log.files.tolist() == ["b.xlsx", "a.xlsx"]
