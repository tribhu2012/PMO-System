"""Spreadsheet import: Excel / CSV workbooks -> validated database tables.

How data arrives in a real PMO
------------------------------
Each data owner keeps one workbook. Jira is assumed to sync into the Jira
workbook (or its CSV export is dropped in), so those sheets use Jira's column
names. Other sheets are maintained by hand by the PMO and HR.

=====================  ==================================  ====================
Workbook               Sheets                              Maintained by
=====================  ==================================  ====================
jira_export.xlsx       Issues, Sprints                     Jira sync / export
pmo_portfolio.xlsx     Projects, Risks, Ceremonies,        PMO, project managers,
                       Archive                             scrum masters
hr_resources.xlsx      Employees, Peer Reviews             HR, team leads
=====================  ==================================  ====================

Sheets refer to each other by *name* (project name, sprint name, assignee),
the way real exports do. Database IDs are assigned during import.

An import is a full snapshot: every table is rebuilt from the files, in one
transaction, and only when validation finds no errors. ``Archive`` is optional;
without it the existing model-training history is kept.

Public API
----------
``read_sources(files)``      -> {sheet: SheetData} plus problems
``build_tables(sheets)``     -> ImportResult (tables + problems + counts)
``validate_files(files)``    -> both steps in one call
``write_workbooks(...)``     -> blank templates or a filled export, as xlsx bytes
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

FIBONACCI_POINTS = (1, 2, 3, 5, 8, 13)


# --------------------------------------------------------------------------- specification
@dataclass(frozen=True)
class Col:
    name: str  # header in the template
    field: str  # internal column name
    kind: str = "text"  # text | int | number | date | choice | percent
    required: bool = True
    choices: tuple = ()
    aliases: dict | None = None  # extra accepted values for a choice -> canonical value
    headers: tuple = ()  # other accepted header spellings
    max_len: int | None = None
    help: str = ""


@dataclass(frozen=True)
class SheetSpec:
    name: str
    workbook: str
    owner: str
    purpose: str
    cols: tuple[Col, ...]
    required: bool = True


JIRA_STATUS = {"open": "To Do", "backlog": "To Do", "selected for development": "To Do",
               "todo": "To Do", "new": "To Do", "in development": "In Progress",
               "doing": "In Progress", "code review": "In Review", "review": "In Review",
               "qa": "In Review", "testing": "In Review", "in qa": "In Review",
               "in test": "In Review", "on hold": "Blocked", "impediment": "Blocked",
               "closed": "Done", "resolved": "Done", "complete": "Done", "completed": "Done"}
JIRA_TYPE = {"user story": "Story", "new feature": "Story", "improvement": "Story",
             "feature": "Story", "sub-task": "Task", "subtask": "Task", "technical task": "Task",
             "defect": "Bug", "research": "Spike"}
JIRA_PRIORITY = {"highest": "Critical", "blocker": "Critical", "major": "High",
                 "minor": "Low", "lowest": "Low", "trivial": "Low"}
SPRINT_STATE = {"complete": "Closed", "completed": "Closed", "done": "Closed",
                "open": "Active", "in progress": "Active", "current": "Active"}
CEREMONY = {"planning": "Sprint Planning", "sprint planning meeting": "Sprint Planning",
            "standup": "Daily Standup", "stand-up": "Daily Standup", "daily scrum": "Daily Standup",
            "daily stand-up": "Daily Standup", "daily": "Daily Standup", "review": "Sprint Review",
            "sprint demo": "Sprint Review", "demo": "Sprint Review", "retro": "Retrospective",
            "sprint retrospective": "Retrospective"}
YES_NO = {"yes": 1, "y": 1, "true": 1, "1": 1, "late": 1, "delayed": 1,
          "no": 0, "n": 0, "false": 0, "0": 0, "on time": 0}

DEPARTMENTS = ("Engineering", "Data Analytics", "Product Management", "Marketing", "Operations")
RISK_CATEGORIES = ("Schedule", "Budget", "Scope", "Technical", "Resource", "Vendor",
                   "Security", "Compliance")

SHEETS: dict[str, SheetSpec] = {s.name: s for s in [
    SheetSpec("Employees", "hr_resources", "HR", "Everyone who can be assigned work.", (
        Col("Employee ID", "employee_id", "int", required=False,
            help="Optional HR number. Leave blank to number people automatically."),
        Col("Employee Name", "employee_name", max_len=100, headers=("name", "full name"),
            help="Must be unique. Other sheets refer to people by this name."),
        Col("Department", "department", max_len=50, help="e.g. " + ", ".join(DEPARTMENTS)),
        Col("Role", "role", max_len=60, help="Project Managers are left out of utilisation."),
        Col("Hire Date", "hire_date", "date"),
        Col("Weekly Capacity (h)", "weekly_capacity_hours", "int",
            headers=("weekly capacity", "capacity hours", "weekly_capacity_hours"),
            help="Contracted hours per week, 1-60."),
    )),
    SheetSpec("Peer Reviews", "hr_resources", "Team leads",
              "Per-sprint peer review score. Optional.", (
        Col("Employee", "employee", headers=("employee name",)),
        Col("Sprint", "sprint", headers=("sprint name",)),
        Col("Peer Review Score", "peer_review_score", "number", help="1-5."),
        Col("Bugs Reopened", "bugs_reopened", "int", required=False,
            help="Bugs reopened against this person's work in the sprint. Blank = 0."),
    ), required=False),
    SheetSpec("Projects", "pmo_portfolio", "PMO", "The project register: one row per project.", (
        Col("Project Name", "project_name", max_len=100, headers=("project", "name")),
        Col("Description", "description", required=False, max_len=200),
        Col("Project Manager", "project_manager", headers=("pm", "manager"),
            help="Name exactly as on the Employees sheet."),
        Col("Department", "department", max_len=50, help="Owning department."),
        Col("Start Date", "start_date", "date"),
        Col("Planned End Date", "planned_end_date", "date", headers=("planned end", "baseline end")),
        Col("Actual End Date", "actual_end_date", "date", required=False,
            help="Fill in only when the project is finished."),
        Col("Budget", "budget", "number", help="Approved budget, in dollars."),
        Col("Actual Cost to Date", "actual_cost", "number",
            headers=("actual cost", "cost to date", "spend to date"),
            help="From Finance. Total spent so far."),
        Col("Status", "status", "choice",
            choices=("On Track", "At Risk", "Delayed", "Completed", "On Hold"),
            aliases={"green": "On Track", "amber": "At Risk", "yellow": "At Risk", "red": "Delayed",
                     "closed": "Completed", "done": "Completed"},
            help="The status the PM reports. The dashboard also computes its own."),
        Col("Priority", "priority", "choice", choices=("Critical", "High", "Medium", "Low")),
    )),
    SheetSpec("Risks", "pmo_portfolio", "Project managers", "The RAID log: one row per risk.", (
        Col("Risk ID", "risk_id", "int", required=False),
        Col("Project", "project", headers=("project name",)),
        Col("Category", "category", max_len=20, help="e.g. " + ", ".join(RISK_CATEGORIES)),
        Col("Description", "risk_description", headers=("risk", "risk description")),
        Col("Probability", "probability", "int", help="1 (rare) to 5 (almost certain)."),
        Col("Impact", "impact", "int", help="1 (minor) to 5 (severe)."),
        Col("Owner", "owner", headers=("risk owner",), help="Name as on the Employees sheet."),
        Col("Status", "status", "choice", choices=("Open", "Mitigating", "Closed", "Accepted"),
            aliases={"in progress": "Mitigating", "mitigation": "Mitigating", "resolved": "Closed",
                     "done": "Closed"}),
        Col("Date Raised", "identified_date", "date", headers=("raised", "identified", "date identified")),
        Col("Date Closed", "resolved_date", "date", required=False,
            headers=("closed", "resolved", "date resolved")),
        Col("Mitigation %", "mitigation_progress", "percent", required=False,
            headers=("mitigation", "mitigation progress"),
            help="How far the mitigation plan has got. 40% or 40. Blank = 0."),
    )),
    SheetSpec("Ceremonies", "pmo_portfolio", "Scrum masters",
              "Planning, stand-ups, reviews and retros. Optional.", (
        Col("Project", "project", headers=("project name",)),
        Col("Sprint", "sprint", required=False, headers=("sprint name",)),
        Col("Ceremony Type", "ceremony_type", "choice",
            choices=("Sprint Planning", "Daily Standup", "Sprint Review", "Retrospective"),
            aliases=CEREMONY, headers=("type", "ceremony")),
        Col("Date", "ceremony_date", "date"),
        Col("Attendance %", "attendance_rate", "percent", headers=("attendance",),
            help="Share of the team who attended. 85% or 85."),
        Col("Duration (min)", "duration_minutes", "int", headers=("duration", "minutes")),
        Col("Action Items", "action_items", "int", required=False, help="Blank = 0."),
    ), required=False),
    SheetSpec("Archive", "pmo_portfolio", "PMO",
              "Closed projects with known outcomes, used to train the delay model. Optional.", (
        Col("Project Code", "project_code", max_len=20),
        Col("Department", "department", max_len=50),
        Col("Close Year", "close_year", "int"),
        Col("Budget Utilization", "budget_utilization", "number",
            help="AC ÷ EV at the mid-project snapshot."),
        Col("Open Risks", "open_risks", "int"),
        Col("Avg Resource Utilization", "avg_resource_utilization", "number", help="e.g. 0.95"),
        Col("Sprint Velocity", "sprint_velocity", "number", help="Average SP of the last 3 sprints."),
        Col("Tasks Remaining", "tasks_remaining", "int"),
        Col("Delayed", "delayed", "choice", choices=("Yes", "No"),
            help="Did the project finish after its planned end date?"),
    ), required=False),
    SheetSpec("Issues", "jira_export", "Jira", "Every Jira issue: stories, tasks, bugs, spikes.", (
        Col("Issue Key", "task_key", max_len=20, headers=("key", "issue id")),
        Col("Summary", "task_name", max_len=200),
        Col("Issue Type", "issue_type", "choice", choices=("Story", "Task", "Bug", "Spike"),
            aliases=JIRA_TYPE, headers=("type",)),
        Col("Project", "project", headers=("project name",),
            help="Project name exactly as on the Projects sheet."),
        Col("Sprint", "sprint", required=False, help="Blank = backlog."),
        Col("Assignee", "assignee", required=False, help="Name as on the Employees sheet."),
        Col("Story Points", "story_points", "int", headers=("story point estimate", "points"),
            help="1, 2, 3, 5, 8 or 13."),
        Col("Original Estimate (h)", "planned_hours", "number",
            headers=("original estimate", "estimate", "planned hours")),
        Col("Time Spent (h)", "actual_hours", "number", required=False,
            headers=("time spent", "actual hours", "logged hours")),
        Col("Status", "status", "choice",
            choices=("To Do", "In Progress", "In Review", "Blocked", "Done"), aliases=JIRA_STATUS),
        Col("Priority", "priority", "choice", choices=("Critical", "High", "Medium", "Low"),
            aliases=JIRA_PRIORITY),
        Col("Created", "created_date", "date", headers=("created date",)),
        Col("Start Date", "started_date", "date", required=False,
            headers=("started", "start", "in progress date"),
            help="When work started. Used for cycle time."),
        Col("Due Date", "due_date", "date", required=False, headers=("due",),
            help="Blank = the sprint's end date."),
        Col("Resolved", "completion_date", "date", required=False,
            headers=("resolution date", "resolved date", "completed", "done date"),
            help="Required when Status is Done."),
    )),
    SheetSpec("Sprints", "jira_export", "Jira", "Every sprint, from Jira's sprint report.", (
        Col("Sprint Name", "sprint_name", max_len=60, headers=("sprint", "name"),
            help="Must be unique. Issues refer to sprints by this name."),
        Col("Project", "project", headers=("project name", "board")),
        Col("Sprint Number", "sprint_number", "int", required=False,
            help="Blank = numbered by start date."),
        Col("Start Date", "start_date", "date"),
        Col("End Date", "end_date", "date"),
        Col("State", "status", "choice", choices=("Closed", "Active"), aliases=SPRINT_STATE,
            headers=("status",)),
        Col("Committed Points", "planned_story_points", "int", required=False,
            headers=("committed", "planned story points", "commitment"),
            help="Blank = summed from the sprint's issues."),
        Col("Completed Points", "completed_story_points", "int", required=False,
            headers=("completed", "completed story points"),
            help="Blank = summed from issues resolved by the sprint end."),
    )),
]}
WORKBOOKS = {
    "jira_export": ("Jira export", ["Issues", "Sprints"]),
    "pmo_portfolio": ("PMO portfolio", ["Projects", "Risks", "Ceremonies", "Archive"]),
    "hr_resources": ("HR resources", ["Employees", "Peer Reviews"]),
}


# --------------------------------------------------------------------------- results
@dataclass
class Problem:
    level: str  # "error" | "warning"
    source: str
    sheet: str
    row: int | None  # spreadsheet row number (header = row 1)
    column: str
    message: str

    def as_dict(self) -> dict:
        return {"Level": self.level.title(), "File": self.source, "Sheet": self.sheet,
                "Row": self.row, "Column": self.column, "Problem": self.message}


@dataclass
class SheetData:
    spec: SheetSpec
    df: pd.DataFrame  # parsed values, internal field names, plus "xrow"
    source: str


@dataclass
class ImportResult:
    tables: dict[str, pd.DataFrame] | None
    problems: list[Problem] = field(default_factory=list)
    sheets_found: dict[str, str] = field(default_factory=dict)  # sheet -> source file

    @property
    def errors(self) -> list[Problem]:
        return [p for p in self.problems if p.level == "error"]

    @property
    def warnings(self) -> list[Problem]:
        return [p for p in self.problems if p.level == "warning"]

    @property
    def ok(self) -> bool:
        return self.tables is not None and not self.errors

    def problem_frame(self) -> pd.DataFrame:
        cols = ["Level", "File", "Sheet", "Row", "Column", "Problem"]
        return pd.DataFrame([p.as_dict() for p in self.problems], columns=cols)


class _Log:
    def __init__(self) -> None:
        self.items: list[Problem] = []

    def error(self, src, sheet, row, col, msg):
        self.items.append(Problem("error", src, sheet, row, col, msg))

    def warn(self, src, sheet, row, col, msg):
        self.items.append(Problem("warning", src, sheet, row, col, msg))


# --------------------------------------------------------------------------- reading
def _norm(s) -> str:
    return re.sub(r"[\s_]+", " ", str(s).replace("*", "")).strip().lower()


def _header_map(spec: SheetSpec, columns) -> dict[str, Col]:
    lookup = {}
    for c in spec.cols:
        for h in (c.name, c.field, *c.headers):
            lookup[_norm(h)] = c
    return {col: lookup[_norm(col)] for col in columns if _norm(col) in lookup}


def _match_score(spec: SheetSpec, columns) -> int:
    found = {c.field for c in _header_map(spec, columns).values()}
    return sum(1 for c in spec.cols if c.required and c.field in found)


def _guess_sheet(columns) -> SheetSpec | None:
    """For CSVs and oddly named tabs: pick the spec whose required headers best match."""
    best, score = None, 0
    for spec in SHEETS.values():
        s = _match_score(spec, columns)
        need = sum(c.required for c in spec.cols)
        if s == need and s > score:
            best, score = spec, s
    return best


def _spec_for_tab(tab: str) -> SheetSpec | None:
    t = _norm(tab)
    for name, spec in SHEETS.items():
        if t in (_norm(name), _norm(name).rstrip("s"), _norm(name.replace(" ", ""))):
            return spec
    aliases = {"raid": "Risks", "raid log": "Risks", "risk register": "Risks", "people": "Employees",
               "staff": "Employees", "register": "Projects", "jira issues": "Issues",
               "project history": "Archive", "history": "Archive", "scrum ceremonies": "Ceremonies"}
    return SHEETS.get(aliases.get(t, ""))


def read_sources(files: list[tuple[str, bytes]]) -> tuple[dict[str, SheetData], list[Problem]]:
    """Read uploaded files. ``files`` is a list of ``(file_name, raw_bytes)``."""
    log = _Log()
    raw: dict[str, tuple[pd.DataFrame, str]] = {}
    for name, data in files:
        frames: list[tuple[str, pd.DataFrame]] = []
        try:
            if name.lower().endswith(".csv"):
                frames = [(Path(name).stem, pd.read_csv(io.BytesIO(data), dtype=object))]
            elif name.lower().endswith((".xlsx", ".xlsm")):
                book = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=object,
                                     engine="openpyxl")
                frames = list(book.items())
            else:
                log.error(name, "", None, "", "Unsupported file type. Upload .xlsx or .csv files.")
                continue
        except Exception as exc:  # corrupt or password-protected file
            log.error(name, "", None, "", f"Could not open the file ({exc.__class__.__name__}: {exc}).")
            continue
        for tab, df in frames:
            if _norm(tab) in ("instructions", "readme", "lists"):
                continue
            df = df.dropna(how="all")
            spec = _spec_for_tab(tab) or _guess_sheet(df.columns)
            if spec is None:
                if len(df):
                    log.warn(name, tab, None, "", "Sheet not recognised, so it was skipped. "
                             "Rename it to one of: " + ", ".join(SHEETS) + ".")
                continue
            if spec.name in raw:
                log.error(name, tab, None, "", f"A {spec.name} sheet was also found in "
                          f"{raw[spec.name][1]}. Upload each sheet once.")
                continue
            raw[spec.name] = (df, name)

    sheets = {}
    for sheet_name, (df, src) in raw.items():
        spec = SHEETS[sheet_name]
        sheets[sheet_name] = _parse_sheet(spec, df, src, log)
    for spec in SHEETS.values():
        if spec.required and spec.name not in raw:
            wb = WORKBOOKS[spec.workbook][0]
            log.error("", spec.name, None, "", f"The {spec.name} sheet is missing. It belongs in "
                      f"the {wb} workbook ({spec.workbook}.xlsx).")
    return sheets, log.items


def _parse_sheet(spec: SheetSpec, df: pd.DataFrame, src: str, log: _Log) -> SheetData:
    hmap = _header_map(spec, df.columns)
    present = {c.field for c in hmap.values()}
    out = pd.DataFrame(index=df.index)
    out["xrow"] = df.index.to_numpy() + 2  # pandas row 0 = spreadsheet row 2
    for c in spec.cols:
        if c.field not in present:
            if c.required:
                log.error(src, spec.name, None, c.name, f"Column '{c.name}' is missing.")
            out[c.field] = np.nan if c.kind != "text" else None
            continue
        header = next(h for h, col in hmap.items() if col.field == c.field)
        out[c.field] = _parse_column(spec, c, df[header], out["xrow"], src, log)
    return SheetData(spec, out.reset_index(drop=True), src)


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and np.isnan(v)) or (isinstance(v, str) and not v.strip()) \
        or v is pd.NaT


def _parse_column(spec, c: Col, s: pd.Series, rows: pd.Series, src, log) -> pd.Series:
    vals = []
    for v, r in zip(s.tolist(), rows.tolist()):
        if _blank(v):
            if c.required:
                log.error(src, spec.name, r, c.name, f"{c.name} is empty.")
            vals.append(None)
            continue
        try:
            vals.append(_convert(c, v))
        except ValueError as exc:
            log.error(src, spec.name, r, c.name, str(exc))
            vals.append(None)
    out = pd.Series(vals, index=s.index, dtype=object)
    if c.kind == "date":
        return pd.to_datetime(out, errors="coerce")
    if c.kind in ("int", "number", "percent"):
        out = pd.to_numeric(out, errors="coerce")
        if c.kind == "percent" and out.notna().any():
            # Excel stores 85% as 0.85; people also type 85. Decide per column.
            out = out * 100 if out.max() <= 1 else out
        return out
    return out


def _convert(c: Col, v):
    if c.kind == "text":
        text = str(v).strip()
        if isinstance(v, float) and v.is_integer():
            text = str(int(v))
        if c.max_len and len(text) > c.max_len:
            raise ValueError(f"{c.name} is {len(text)} characters; the limit is {c.max_len}.")
        return text
    if c.kind in ("int", "number", "percent"):
        if isinstance(v, str):
            v = v.strip().replace(",", "").replace("$", "")
            if c.kind == "percent" and v.endswith("%"):
                v = float(v[:-1]) / 100
        try:
            x = float(v)
        except (TypeError, ValueError):
            raise ValueError(f"{c.name} must be a number, got '{v}'.") from None
        if c.kind == "int" and not float(x).is_integer():
            raise ValueError(f"{c.name} must be a whole number, got {v}.")
        return int(x) if c.kind == "int" else x
    if c.kind == "date":
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            d = pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))  # Excel serial date
        else:
            d = pd.to_datetime(str(v).strip(), errors="coerce")
        if pd.isna(d):
            raise ValueError(f"{c.name} is not a date: '{v}'. Use YYYY-MM-DD.")
        return pd.Timestamp(d).normalize()
    if c.kind == "choice":
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        key = _norm(v)
        canon = {_norm(x): x for x in c.choices}
        if c.field == "delayed":
            if key in YES_NO:
                return YES_NO[key]
        elif key in canon:
            return canon[key]
        elif c.aliases and key in c.aliases:
            return c.aliases[key]
        raise ValueError(f"'{v}' is not a valid {c.name}. Use one of: {', '.join(c.choices)}.")
    raise AssertionError(c.kind)


# --------------------------------------------------------------------------- building
def validate_files(files: list[tuple[str, bytes]], existing_history: pd.DataFrame | None = None
                   ) -> ImportResult:
    sheets, problems = read_sources(files)
    result = build_tables(sheets, existing_history)
    result.problems = problems + result.problems
    if any(p.level == "error" for p in result.problems):
        result.tables = None
    return result


def build_tables(sheets: dict[str, SheetData], existing_history: pd.DataFrame | None = None
                 ) -> ImportResult:
    """Cross-check the parsed sheets and assemble database-ready tables."""
    log = _Log()
    found = {k: v.source for k, v in sheets.items()}
    if any(name not in sheets for name, spec in SHEETS.items() if spec.required):
        return ImportResult(None, log.items, found)

    employees = _employees(sheets["Employees"], log)
    people = _lookup(employees.employee_name, employees.employee_id)
    projects = _projects(sheets["Projects"], people, log)
    proj_ids = _lookup(projects.project_name, projects.project_id)
    sprints = _sprints(sheets["Sprints"], proj_ids, log)
    tasks = _tasks(sheets["Issues"], proj_ids, sprints, people, log)
    as_of = _as_of(tasks, sheets.get("Ceremonies"))
    sprints = _sprint_points(sprints, tasks, as_of, sheets["Sprints"], log)
    risks = _risks(sheets["Risks"], proj_ids, people, log)
    ceremonies = (_ceremonies(sheets["Ceremonies"], proj_ids, sprints, log) if "Ceremonies" in sheets
                  else _empty_ceremonies())
    if "Ceremonies" not in sheets:
        log.warn("", "Ceremonies", None, "", "No Ceremonies sheet, so attendance charts will be empty.")
    performance = _performance(tasks, sprints, as_of, sheets.get("Peer Reviews"), people, log)
    if "Archive" in sheets:
        history = _history(sheets["Archive"], log)
    elif existing_history is not None and len(existing_history):
        history = existing_history.copy()
    else:
        log.error("", "Archive", None, "", "No Archive sheet and no saved history to keep. The "
                  "delay model needs closed projects with known outcomes to learn from.")
        history = None

    if any(p.level == "error" for p in log.items) or history is None:
        return ImportResult(None, log.items, found)
    tables = {
        "employees": employees, "projects": projects,
        "sprints": sprints.drop(columns=["_key"]), "tasks": tasks,
        "risks": risks, "scrum_ceremonies": ceremonies,
        "performance": performance, "project_history": history,
    }
    return ImportResult(tables, log.items, found)


def _key(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def _lookup(names: pd.Series, ids: pd.Series) -> dict[str, int]:
    return dict(zip(names.map(_key), ids))


def _resolve(sd: SheetData, col: str, header: str, lookup: dict, log: _Log, what: str,
             required: bool = True) -> pd.Series:
    out = []
    for v, r in zip(sd.df[col], sd.df.xrow):
        if v is None or (isinstance(v, float) and np.isnan(v)):
            out.append(None)
            continue
        hit = lookup.get(_key(v))
        if hit is None:
            log.error(sd.source, sd.spec.name, r, header, f"'{v}' is not on the {what} sheet. "
                      "Check the spelling or add it there first.")
        out.append(hit)
    return pd.Series(out, index=sd.df.index, dtype=object)


def _dupes(sd: SheetData, col: str, header: str, log: _Log, keyfn=_key) -> None:
    keys = sd.df[col].map(lambda v: keyfn(v) if v is not None else None)
    dup = keys.notna() & keys.duplicated(keep="first")
    for v, r in zip(sd.df.loc[dup, col], sd.df.loc[dup, "xrow"]):
        log.error(sd.source, sd.spec.name, r, header, f"'{v}' appears more than once. "
                  f"Each {header} must be unique.")


def _ids(sd: SheetData, col: str) -> pd.Series:
    given = sd.df[col]
    if given.notna().all() and not given.duplicated().any() and len(given):
        return given.astype(int)
    return pd.Series(np.arange(1, len(sd.df) + 1), index=sd.df.index)


def _employees(sd: SheetData, log: _Log) -> pd.DataFrame:
    d = sd.df
    _dupes(sd, "employee_name", "Employee Name", log)
    for v, r in zip(d.weekly_capacity_hours, d.xrow):
        if pd.notna(v) and not 1 <= v <= 60:
            log.error(sd.source, sd.spec.name, r, "Weekly Capacity (h)", f"{v:g} hours is outside 1-60.")
    if d.employee_id.notna().any() and (d.employee_id.isna().any() or d.employee_id.duplicated().any()):
        log.warn(sd.source, sd.spec.name, None, "Employee ID", "Some IDs are blank or repeated, so "
                 "everyone was renumbered 1, 2, 3...")
    return pd.DataFrame({
        "employee_id": _ids(sd, "employee_id"), "employee_name": d.employee_name,
        "department": d.department, "role": d.role, "hire_date": d.hire_date,
        "weekly_capacity_hours": d.weekly_capacity_hours,
    })


def _projects(sd: SheetData, people: dict, log: _Log) -> pd.DataFrame:
    d = sd.df
    _dupes(sd, "project_name", "Project Name", log)
    pm = _resolve(sd, "project_manager", "Project Manager", people, log, "Employees")
    src, sh = sd.source, sd.spec.name
    for r in d.itertuples():
        if pd.notna(r.start_date) and pd.notna(r.planned_end_date) and r.planned_end_date <= r.start_date:
            log.error(src, sh, r.xrow, "Planned End Date", "Planned end must be after the start date.")
        if pd.notna(r.actual_end_date) and pd.notna(r.start_date) and r.actual_end_date <= r.start_date:
            log.error(src, sh, r.xrow, "Actual End Date", "Actual end must be after the start date.")
        if pd.notna(r.budget) and r.budget <= 0:
            log.error(src, sh, r.xrow, "Budget", "Budget must be greater than 0.")
        if pd.notna(r.actual_cost) and r.actual_cost < 0:
            log.error(src, sh, r.xrow, "Actual Cost to Date", "Cost cannot be negative.")
        if r.status == "Completed" and pd.isna(r.actual_end_date):
            log.error(src, sh, r.xrow, "Actual End Date", "Status is Completed but there is no actual end date.")
        if pd.notna(r.actual_end_date) and r.status not in ("Completed", None):
            log.warn(src, sh, r.xrow, "Status", f"Has an actual end date but status is {r.status}. "
                     "The dashboard treats it as completed.")
    return pd.DataFrame({
        "project_id": np.arange(1, len(d) + 1), "project_name": d.project_name,
        "description": d.description, "project_manager_id": pm, "department": d.department,
        "start_date": d.start_date, "planned_end_date": d.planned_end_date,
        "actual_end_date": d.actual_end_date, "budget": d.budget, "actual_cost": d.actual_cost,
        "status": d.status, "priority": d.priority,
    })


def _sprints(sd: SheetData, proj_ids: dict, log: _Log) -> pd.DataFrame:
    d = sd.df
    _dupes(sd, "sprint_name", "Sprint Name", log)
    pid = _resolve(sd, "project", "Project", proj_ids, log, "Projects")
    for r in d.itertuples():
        if pd.notna(r.start_date) and pd.notna(r.end_date) and r.end_date <= r.start_date:
            log.error(sd.source, sd.spec.name, r.xrow, "End Date", "End date must be after the start date.")
    s = pd.DataFrame({"project_id": pid, "sprint_name": d.sprint_name, "start_date": d.start_date,
                      "end_date": d.end_date, "status": d.status, "sprint_number": d.sprint_number,
                      "planned_story_points": d.planned_story_points,
                      "completed_story_points": d.completed_story_points, "xrow": d.xrow})
    s["_auto"] = np.nan
    ok = s.project_id.notna() & s.start_date.notna()
    s.loc[ok, "_auto"] = s[ok].groupby("project_id").start_date.rank(method="first")
    s["sprint_number"] = s.sprint_number.fillna(s.get("_auto"))
    dup = s.dropna(subset=["project_id", "sprint_number"]).duplicated(["project_id", "sprint_number"])
    for r in s.loc[dup[dup].index].itertuples():
        log.error(sd.source, sd.spec.name, r.xrow, "Sprint Number",
                  f"Sprint number {int(r.sprint_number)} is used twice in the same project.")
    s = s.sort_values(["start_date", "project_id"], kind="stable").reset_index(drop=True)
    s.insert(0, "sprint_id", np.arange(1, len(s) + 1))
    s["_key"] = s.sprint_name.map(_key)
    return s.drop(columns=["_auto"], errors="ignore")


def _tasks(sd: SheetData, proj_ids: dict, sprints: pd.DataFrame, people: dict, log: _Log
           ) -> pd.DataFrame:
    d = sd.df
    src, sh = sd.source, sd.spec.name
    _dupes(sd, "task_key", "Issue Key", log, keyfn=lambda v: str(v).strip().upper())
    pid = _resolve(sd, "project", "Project", proj_ids, log, "Projects")
    sprint_by_name = dict(zip(sprints._key, sprints.sprint_id))
    sid = _resolve(sd, "sprint", "Sprint", sprint_by_name, log, "Sprints")
    eid = _resolve(sd, "assignee", "Assignee", people, log, "Employees")
    sp_project = sprints.set_index("sprint_id").project_id
    sp_end = sprints.set_index("sprint_id").end_date
    due = d.due_date.copy()
    for i, r in enumerate(d.itertuples()):
        s = sid.iloc[i]
        if s is not None and pid.iloc[i] is not None and sp_project.get(s) != pid.iloc[i]:
            log.error(src, sh, r.xrow, "Sprint", f"Sprint '{r.sprint}' belongs to a different project.")
        if pd.notna(r.story_points) and int(r.story_points) not in FIBONACCI_POINTS:
            log.error(src, sh, r.xrow, "Story Points", f"{int(r.story_points)} is not a valid estimate. "
                      "Use 1, 2, 3, 5, 8 or 13.")
        if pd.notna(r.planned_hours) and r.planned_hours < 0:
            log.error(src, sh, r.xrow, "Original Estimate (h)", "Hours cannot be negative.")
        if r.status == "Done" and pd.isna(r.completion_date):
            log.error(src, sh, r.xrow, "Resolved", "Status is Done but there is no Resolved date.")
        if pd.notna(r.completion_date) and r.status not in ("Done", None):
            log.warn(src, sh, r.xrow, "Resolved", f"Has a Resolved date but status is {r.status}. "
                     "It is counted as done on that date.")
        if pd.notna(r.completion_date) and pd.notna(r.created_date) and r.completion_date < r.created_date:
            log.error(src, sh, r.xrow, "Resolved", "Resolved date is before the Created date.")
        if pd.isna(r.due_date):
            if s is not None and s in sp_end.index:
                due.iloc[i] = sp_end[s]
            elif pd.notna(r.created_date):
                log.error(src, sh, r.xrow, "Due Date", "Due date is empty and the issue is not in a sprint.")
    return pd.DataFrame({
        "task_id": np.arange(1, len(d) + 1), "task_key": d.task_key.map(lambda v: str(v).strip().upper()
                                                                        if v is not None else v),
        "project_id": pid, "sprint_id": sid, "employee_id": eid, "task_name": d.task_name,
        "issue_type": d.issue_type, "story_points": d.story_points, "planned_hours": d.planned_hours,
        "actual_hours": d.actual_hours, "status": d.status, "priority": d.priority,
        "created_date": d.created_date, "started_date": d.started_date,
        "due_date": pd.to_datetime(due), "completion_date": d.completion_date,
    })


def _as_of(tasks: pd.DataFrame, ceremonies: SheetData | None) -> pd.Timestamp:
    dates = [tasks.completion_date.max(), tasks.created_date.max()]
    if ceremonies is not None:
        dates.append(ceremonies.df.ceremony_date.max())
    return pd.Timestamp(max(d for d in dates if pd.notna(d))).normalize()


def _sprint_points(sprints, tasks, as_of, sd: SheetData, log: _Log) -> pd.DataFrame:
    """Fill blank committed / completed points from the issues (same rule as the generator)."""
    t = tasks[tasks.sprint_id.notna()].merge(sprints[["sprint_id", "end_date"]], on="sprint_id")
    cutoff = t.end_date.where(t.end_date <= as_of, as_of)
    t["sp_done"] = t.story_points.where(t.completion_date.notna() & (t.completion_date <= cutoff), 0)
    agg = t.groupby("sprint_id").agg(planned=("story_points", "sum"), done=("sp_done", "sum"))
    s = sprints.copy()
    s["planned_story_points"] = s.planned_story_points.fillna(s.sprint_id.map(agg.planned)).fillna(0)
    s["completed_story_points"] = s.completed_story_points.fillna(s.sprint_id.map(agg.done)).fillna(0)
    empty = ~s.sprint_id.isin(agg.index)
    for r in s[empty].itertuples():
        log.warn(sd.source, sd.spec.name, r.xrow, "Sprint Name", f"No issues are assigned to '{r.sprint_name}'.")
    cols = ["sprint_id", "project_id", "sprint_number", "sprint_name", "start_date", "end_date",
            "status", "planned_story_points", "completed_story_points", "_key"]
    return s[cols]


def _risks(sd: SheetData, proj_ids: dict, people: dict, log: _Log) -> pd.DataFrame:
    d = sd.df
    src, sh = sd.source, sd.spec.name
    pid = _resolve(sd, "project", "Project", proj_ids, log, "Projects")
    owner = _resolve(sd, "owner", "Owner", people, log, "Employees")
    for r in d.itertuples():
        for col, header in (("probability", "Probability"), ("impact", "Impact")):
            v = getattr(r, col)
            if pd.notna(v) and not 1 <= v <= 5:
                log.error(src, sh, r.xrow, header, f"{header} must be 1-5, got {v:g}.")
        if pd.notna(r.mitigation_progress) and not 0 <= r.mitigation_progress <= 100:
            log.error(src, sh, r.xrow, "Mitigation %", "Mitigation must be between 0% and 100%.")
        if r.status == "Closed" and pd.isna(r.resolved_date):
            log.error(src, sh, r.xrow, "Date Closed", "Status is Closed but there is no Date Closed.")
        if pd.notna(r.resolved_date) and pd.notna(r.identified_date) and r.resolved_date < r.identified_date:
            log.error(src, sh, r.xrow, "Date Closed", "Date Closed is before Date Raised.")
    return pd.DataFrame({
        "risk_id": _ids(sd, "risk_id"), "project_id": pid, "category": d.category,
        "risk_description": d.risk_description, "probability": d.probability, "impact": d.impact,
        "risk_score": d.probability * d.impact, "owner_id": owner, "status": d.status,
        "identified_date": d.identified_date, "resolved_date": d.resolved_date,
        "mitigation_progress": d.mitigation_progress.fillna(0).round(),
    })


def _ceremonies(sd: SheetData, proj_ids: dict, sprints: pd.DataFrame, log: _Log) -> pd.DataFrame:
    d = sd.df
    pid = _resolve(sd, "project", "Project", proj_ids, log, "Projects")
    sid = _resolve(sd, "sprint", "Sprint", dict(zip(sprints._key, sprints.sprint_id)), log, "Sprints")
    for r in d.itertuples():
        if pd.notna(r.attendance_rate) and not 0 <= r.attendance_rate <= 100:
            log.error(sd.source, sd.spec.name, r.xrow, "Attendance %", "Attendance must be 0%-100%.")
    return pd.DataFrame({
        "ceremony_id": np.arange(1, len(d) + 1), "project_id": pid, "sprint_id": sid,
        "ceremony_type": d.ceremony_type, "ceremony_date": d.ceremony_date,
        "attendance_rate": (d.attendance_rate / 100).round(3), "duration_minutes": d.duration_minutes,
        "action_items": d.action_items.fillna(0),
    })


def _empty_ceremonies() -> pd.DataFrame:
    return pd.DataFrame(columns=["ceremony_id", "project_id", "sprint_id", "ceremony_type",
                                 "ceremony_date", "attendance_rate", "duration_minutes", "action_items"])


def _performance(tasks, sprints, as_of, reviews: SheetData | None, people: dict, log: _Log
                 ) -> pd.DataFrame:
    """One row per person per sprint, counted from the issues; peer reviews joined on."""
    t = tasks[tasks.sprint_id.notna() & tasks.employee_id.notna()].merge(
        sprints[["sprint_id", "end_date"]], on="sprint_id")
    cutoff = t.end_date.where(t.end_date <= as_of, as_of)
    t["done"] = t.completion_date.notna() & (t.completion_date <= cutoff)
    t["sp_done"] = t.story_points.where(t.done, 0)
    g = t.groupby(["employee_id", "sprint_id"]).agg(
        stories_assigned=("task_id", "count"), stories_completed=("done", "sum"),
        story_points_completed=("sp_done", "sum")).reset_index()
    g["bugs_reopened"] = 0
    g["peer_review_score"] = np.nan
    if reviews is not None:
        d = reviews.df
        e = _resolve(reviews, "employee", "Employee", people, log, "Employees")
        s = _resolve(reviews, "sprint", "Sprint", dict(zip(sprints._key, sprints.sprint_id)), log, "Sprints")
        rv = pd.DataFrame({"employee_id": e, "sprint_id": s, "score": d.peer_review_score,
                           "bugs": d.bugs_reopened, "xrow": d.xrow}).dropna(subset=["employee_id", "sprint_id"])
        for r in rv.itertuples():
            if pd.notna(r.score) and not 1 <= r.score <= 5:
                log.error(reviews.source, "Peer Reviews", r.xrow, "Peer Review Score", "Score must be 1-5.")
        dup = rv.duplicated(["employee_id", "sprint_id"])
        for r in rv[dup].itertuples():
            log.error(reviews.source, "Peer Reviews", r.xrow, "Sprint", "This person already has a review "
                      "for this sprint.")
        rv = rv[~dup].astype({"employee_id": int, "sprint_id": int})
        merged = g.merge(rv, on=["employee_id", "sprint_id"], how="outer", indicator="origin")
        for r in merged[merged.origin == "right_only"].itertuples():
            log.warn(reviews.source, "Peer Reviews", r.xrow, "Sprint",
                     "This person has no issues in that sprint, so the review was not used.")
        g = merged[merged.origin != "right_only"].copy()
        g["peer_review_score"] = g.score.round(1)
        g["bugs_reopened"] = g.bugs.fillna(0)
        g = g.drop(columns=["score", "bugs", "xrow", "origin"])
    missing = int(g.peer_review_score.isna().sum())
    if missing:
        log.warn(reviews.source if reviews else "", "Peer Reviews", None, "Peer Review Score",
                 f"{missing:,} person-sprints have no peer review. Their performance score uses the "
                 "other three components.")
    for c in ("stories_assigned", "stories_completed", "story_points_completed"):
        g[c] = g[c].astype(int)
    g["bugs_reopened"] = np.minimum(g.bugs_reopened, g.stories_completed).astype(int)
    cols = ["employee_id", "sprint_id", "stories_assigned", "stories_completed",
            "story_points_completed", "bugs_reopened", "peer_review_score"]
    return g[cols].astype({"employee_id": int, "sprint_id": int}).sort_values(
        ["sprint_id", "employee_id"]).reset_index(drop=True)


def _history(sd: SheetData, log: _Log) -> pd.DataFrame:
    d = sd.df
    if d.delayed.nunique() < 2:
        log.error(sd.source, sd.spec.name, None, "Delayed", "The archive needs both delayed and on-time "
                  "projects, or the model cannot learn the difference.")
    if len(d) < 40:
        log.warn(sd.source, sd.spec.name, None, "", f"Only {len(d)} archived projects. Model accuracy "
                 "will be unreliable below about 100.")
    return pd.DataFrame({"history_id": np.arange(1, len(d) + 1), **{
        c.field: d[c.field] for c in sd.spec.cols}})


# --------------------------------------------------------------------------- writing
def tables_to_sheets(tables: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """Database tables -> sheet rows (headers as in the templates). Used for exports and samples."""
    e, p, s, t = tables["employees"], tables["projects"], tables["sprints"], tables["tasks"]
    name = e.set_index("employee_id").employee_name
    pname = p.set_index("project_id").project_name
    sname = s.set_index("sprint_id").sprint_name

    def m(series, mapping):
        return series.map(lambda v: mapping.get(int(v)) if pd.notna(v) else None)

    out = {
        "Employees": pd.DataFrame({
            "Employee ID": e.employee_id, "Employee Name": e.employee_name, "Department": e.department,
            "Role": e.role, "Hire Date": e.hire_date, "Weekly Capacity (h)": e.weekly_capacity_hours}),
        "Projects": pd.DataFrame({
            "Project Name": p.project_name, "Description": p.description,
            "Project Manager": m(p.project_manager_id, name), "Department": p.department,
            "Start Date": p.start_date, "Planned End Date": p.planned_end_date,
            "Actual End Date": p.actual_end_date, "Budget": p.budget,
            "Actual Cost to Date": p.actual_cost, "Status": p.status, "Priority": p.priority}),
        "Sprints": pd.DataFrame({
            "Sprint Name": s.sprint_name, "Project": m(s.project_id, pname),
            "Sprint Number": s.sprint_number, "Start Date": s.start_date, "End Date": s.end_date,
            "State": s.status, "Committed Points": s.planned_story_points,
            "Completed Points": s.completed_story_points}),
        "Issues": pd.DataFrame({
            "Issue Key": t.task_key, "Summary": t.task_name, "Issue Type": t.issue_type,
            "Project": m(t.project_id, pname), "Sprint": m(t.sprint_id, sname),
            "Assignee": m(t.employee_id, name), "Story Points": t.story_points,
            "Original Estimate (h)": t.planned_hours, "Time Spent (h)": t.actual_hours,
            "Status": t.status, "Priority": t.priority, "Created": t.created_date,
            "Start Date": t.started_date, "Due Date": t.due_date, "Resolved": t.completion_date}),
    }
    r = tables["risks"]
    out["Risks"] = pd.DataFrame({
        "Risk ID": r.risk_id, "Project": m(r.project_id, pname), "Category": r.category,
        "Description": r.risk_description, "Probability": r.probability, "Impact": r.impact,
        "Owner": m(r.owner_id, name), "Status": r.status, "Date Raised": r.identified_date,
        "Date Closed": r.resolved_date, "Mitigation %": r.mitigation_progress / 100})
    c = tables["scrum_ceremonies"]
    out["Ceremonies"] = pd.DataFrame({
        "Project": m(c.project_id, pname), "Sprint": m(c.sprint_id, sname),
        "Ceremony Type": c.ceremony_type, "Date": c.ceremony_date,
        "Attendance %": c.attendance_rate, "Duration (min)": c.duration_minutes,
        "Action Items": c.action_items})
    f = tables["performance"]
    out["Peer Reviews"] = pd.DataFrame({
        "Employee": m(f.employee_id, name), "Sprint": m(f.sprint_id, sname),
        "Peer Review Score": f.peer_review_score, "Bugs Reopened": f.bugs_reopened})
    h = tables["project_history"]
    out["Archive"] = pd.DataFrame({
        "Project Code": h.project_code, "Department": h.department, "Close Year": h.close_year,
        "Budget Utilization": h.budget_utilization, "Open Risks": h.open_risks,
        "Avg Resource Utilization": h.avg_resource_utilization, "Sprint Velocity": h.sprint_velocity,
        "Tasks Remaining": h.tasks_remaining, "Delayed": h.delayed.map({1: "Yes", 0: "No"})})
    return out


def write_workbooks(tables: dict[str, pd.DataFrame] | None = None) -> dict[str, bytes]:
    """Build the three workbooks as xlsx bytes: blank templates, or filled with ``tables``."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = tables_to_sheets(tables) if tables is not None else {}
    head_req = PatternFill("solid", fgColor="1F5FAF")
    head_opt = PatternFill("solid", fgColor="7FA6D6")
    out = {}
    for wb_key, (wb_title, sheet_names) in WORKBOOKS.items():
        wb = Workbook()
        info = wb.active
        info.title = "Instructions"
        info.append([f"{wb_title} workbook - ProjectPulse import"])
        info["A1"].font = Font(bold=True, size=14)
        info.append(["Fill one row per record. Dark blue headers are required, light blue are optional. "
                     "Keep the sheet names and headers as they are. Dates: YYYY-MM-DD."])
        info.append([])
        info.append(["Sheet", "Column", "Required", "Allowed values / notes"])
        for cell in info[4]:
            cell.font = Font(bold=True)
        for sname in sheet_names:
            spec = SHEETS[sname]
            info.append([sname, "", "", f"{spec.purpose} Maintained by: {spec.owner}."
                         + ("" if spec.required else " (whole sheet optional)")])
            info.cell(info.max_row, 1).font = Font(bold=True)
            for c in spec.cols:
                note = ", ".join(c.choices) if c.choices else c.help
                if c.choices and c.help:
                    note += f". {c.help}"
                info.append(["", c.name, "Yes" if c.required else "No", note])
        for col, w in zip("ABCD", (16, 26, 10, 90)):
            info.column_dimensions[col].width = w

        for sname in sheet_names:
            spec = SHEETS[sname]
            ws = wb.create_sheet(sname)
            ws.append([c.name for c in spec.cols])
            for i, c in enumerate(spec.cols, start=1):
                cell = ws.cell(1, i)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_req if c.required else head_opt
                cell.alignment = Alignment(vertical="center")
                if c.help or c.choices:
                    cell.comment = Comment(c.help or ", ".join(c.choices), "ProjectPulse")
                ws.column_dimensions[get_column_letter(i)].width = max(14, min(42, len(c.name) + 6))
            ws.freeze_panes = "A2"
            df = rows.get(sname)
            n = 0 if df is None else len(df)
            if df is not None:
                for rec in df.itertuples(index=False):
                    ws.append([_cell(v) for v in rec])
            last = max(n + 1, 2) + 500
            for i, c in enumerate(spec.cols, start=1):
                letter = get_column_letter(i)
                if c.choices:
                    dv = DataValidation(type="list", formula1='"' + ",".join(c.choices) + '"',
                                        allow_blank=not c.required, showErrorMessage=False)
                    ws.add_data_validation(dv)
                    dv.add(f"{letter}2:{letter}{last}")
                fmt = {"date": "yyyy-mm-dd", "percent": "0%", "number": "#,##0.00"}.get(c.kind)
                if c.field in ("budget", "actual_cost"):
                    fmt = "$#,##0"
                if fmt:
                    for row in ws.iter_rows(min_row=2, max_row=last, min_col=i, max_col=i):
                        row[0].number_format = fmt
        buf = io.BytesIO()
        wb.save(buf)
        out[f"{wb_key}.xlsx"] = buf.getvalue()
    return out


def _cell(v):
    if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().date()
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    return v
